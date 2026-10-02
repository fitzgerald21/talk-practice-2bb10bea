"""Builds practice/index.html (the flashcard app) and the Word script doc from script_data.py.
Run:  python3 build.py
"""
import json, os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from script_data import SLIDES

here = os.path.dirname(os.path.abspath(__file__))
root = os.path.dirname(here)

# ---- app ----
data = [{"title": s["title"], "time": s["time"],
         "chunks": [{"cue": c[0], "kind": c[1], "text": c[2]} for c in s["chunks"]]} for s in SLIDES]
html = open(os.path.join(here, "template.html"), encoding="utf-8").read()
html = html.replace("__DATA__", json.dumps(data, ensure_ascii=False)).replace("__VERSION__", "v5-1")
open(os.path.join(here, "index.html"), "w", encoding="utf-8").write(html)

# ---- word doc ----
d = Document()
for s in d.sections:
    s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(0.9)
d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(11)
d.add_heading("Getting Results Without Authority: Speaking Script (v5)", 0)
p = d.add_paragraph("What to say, and when. Each slide has a target time. Bold labels are the cue for each beat; ")
r = p.add_run("[bracketed orange text]"); r.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
p.add_run(" is something to do rather than say.")

t = d.add_table(rows=1, cols=3); t.style = "Light Grid Accent 1"
for i, h in enumerate(["#", "Slide", "Time"]): t.rows[0].cells[i].text = h
for i, s in enumerate(SLIDES, 1):
    row = t.add_row().cells
    row[0].text = str(i); row[1].text = s["title"]; row[2].text = s["time"]
d.add_paragraph()
n = d.add_paragraph(); n.add_run("Note: slide 7's 4:00 happens in breakouts. Speaking time is about 8.5 minutes, plus the 4-minute breakout.").italic = True

for i, s in enumerate(SLIDES, 1):
    d.add_heading(f"Slide {i}: {s['title']}", 1)
    tp = d.add_paragraph(); tr = tp.add_run(f"Time: {s['time']}"); tr.italic = True
    for cue, kind, text in s["chunks"]:
        para = d.add_paragraph(style="List Bullet")
        para.add_run(cue + ": ").bold = True
        run = para.add_run(text if kind == "say" else f"[{text.strip('[]')}]")
        if kind == "do": run.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
d.save(os.path.join(root, "v5 Speaking Script.docx"))
print("built")
